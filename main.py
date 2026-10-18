"""Парсер курсов валют ЦБ РФ с выгрузкой в Excel."""
import json
from datetime import datetime

import requests
import pandas as pd

# Отключаем системный прокси (v2rayN)
session = requests.Session()
session.trust_env = False

URL = "https://www.cbr-xml-daily.ru/daily_json.js"


def get_rates() -> dict:
    """Загружает курсы валют с ЦБ РФ."""
    response = session.get(URL, timeout=15)
    response.raise_for_status()
    return response.json()


def parse_rates(data: dict) -> pd.DataFrame:
    """Преобразует JSON в DataFrame."""
    date = datetime.fromisoformat(data["Date"]).strftime("%d.%m.%Y")
    rows = []

    for code, info in data["Valute"].items():
        nominal = info.get("Nominal", 1)
        value = info["Value"] / nominal
        previous = info["Previous"] / nominal
        change = value - previous
        change_pct = (change / previous) * 100 if previous else 0

        rows.append({
            "Код": code,
            "Название": info["Name"],
            "Курс (₽)": round(value, 4),
            "Было (₽)": round(previous, 4),
            "Изменение (₽)": round(change, 4),
            "Изменение (%)": round(change_pct, 2),
            "Дата": date,
        })

    df = pd.DataFrame(rows)
    df = df.sort_values("Код").reset_index(drop=True)
    return df


def save_to_excel(df: pd.DataFrame, filename: str = "rates.xlsx") -> None:
    """Сохраняет DataFrame в Excel с оформлением."""
    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Курсы валют")

        from openpyxl.styles import Font, PatternFill

        # ... внутри with pd.ExcelWriter:

        worksheet = writer.sheets["Курсы валют"]

        # Заголовки — жирные, с серым фоном
        header_fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill

        # Цвет для изменения (%)
        green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

        change_col = None
        for col_idx, cell in enumerate(worksheet[1], start=1):
            if cell.value == "Изменение (%)":
                change_col = col_idx
                break

        if change_col:
            for row in worksheet.iter_rows(min_row=2, min_col=change_col, max_col=change_col):
                cell = row[0]
                if isinstance(cell.value, (int, float)):
                    if cell.value > 0:
                        cell.fill = green_fill
                    elif cell.value < 0:
                        cell.fill = red_fill
        worksheet = writer.sheets["Курсы валют"]
        # Автоширина колонок
        for col in worksheet.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            worksheet.column_dimensions[col[0].column_letter].width = max_len + 3

        # Жирные заголовки
        from openpyxl.styles import Font
        for cell in worksheet[1]:
            cell.font = Font(bold=True)


if __name__ == "__main__":
    print("Загружаю курсы валют...")
    data = get_rates()
    df = parse_rates(data)

    print(f"Дата: {df['Дата'].iloc[0]}")
    print(f"Валют: {len(df)}")
    print("\nПервые 5 валют:")
    print(df.head()[["Код", "Название", "Курс (₽)", "Изменение (%)"]].to_string(index=False))

    save_to_excel(df)
    print(f"\n✅ Сохранено в rates.xlsx")